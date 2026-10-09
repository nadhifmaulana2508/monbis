from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Monbis_Dokumentasi_Report_Atasan.docx"

NAVY = "17365D"
BLUE = "1F4E79"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
MID_BLUE = "D9EAF4"
TEAL = "008C95"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
GRAY = "5B6573"
LIGHT_GRAY = "F3F6F8"
BORDER = "C9D8E4"
WHITE = "FFFFFF"
BLACK = "000000"


REPORT_GROUPS = [
    {
        "name": "Ringkasan Eksekutif",
        "audience": "Direksi, manajemen, dan pimpinan unit",
        "objective": "Memberikan gambaran cepat mengenai posisi bisnis, kualitas aset, tren kinerja, dan area yang membutuhkan perhatian.",
        "reports": [
            ("dashboard", "Dashboard", "Ringkasan utama kinerja Monbis dalam satu halaman.", "KPI bisnis, tren, alert, dan drill-down ke laporan terkait.", "Manajemen"),
        ],
    },
    {
        "name": "Pemasaran",
        "audience": "Pimpinan bisnis, kantor cabang, dan AO",
        "objective": "Memantau produksi kredit, target pertumbuhan, pipeline, refinancing, dan perubahan kualitas sejak proses pemasaran sampai realisasi.",
        "reports": [
            ("realisasi_kredit_growth", "Realisasi Kredit", "Memantau realisasi kredit dan pertumbuhannya menurut periode serta unit kerja.", "Nominal realisasi, pertumbuhan, periode, kantor, dan detail transaksi.", "Manajemen bisnis"),
            ("realisasi_ao", "Realisasi Kredit AO", "Membaca kontribusi produksi dan kualitas portofolio yang dikelola setiap AO.", "Produksi AO, kelolaan, flow, dan rekening pembentuk.", "Pimpinan cabang / AO"),
            ("realisasi_rbb", "Produksi vs RBB", "Membandingkan produksi aktual dengan target RBB bulan berjalan dan target tahunan.", "Target, realisasi, pencapaian, selisih, kekurangan, dan tren.", "Manajemen bisnis"),
            ("migrasi_bucket_sc", "Migrasi Bucket SC", "Mengikuti perubahan bucket kualitas debitur pada proses monitoring dan collection.", "Perpindahan bucket, nominal, NOA, dan rekening pembentuk.", "Collection / bisnis"),
            ("pipelane_ao_jt", "Pipeline AO Kredit", "Memantau peluang kredit yang masih berada di pipeline AO sampai jatuh tempo proses.", "Daftar pipeline, AO, nilai peluang, status, dan tanggal penting.", "Pimpinan bisnis / AO"),
            ("jatuh_tempo", "Jatuh Tempo & Refinancing", "Menyiapkan tindak lanjut terhadap kredit yang mendekati jatuh tempo dan peluang refinancing.", "Rekening, debitur, tanggal jatuh tempo, outstanding, dan status tindak lanjut.", "Bisnis / collection"),
        ],
    },
    {
        "name": "Monitoring Pembayaran",
        "audience": "Pimpinan cabang, collection, dan AO",
        "objective": "Menilai ketepatan pembayaran, repayment rate, potensi penurunan kualitas, serta posisi debitur yang perlu segera ditindaklanjuti.",
        "reports": [
            ("search_debitur", "Search Debitur Kredit", "Mencari profil dan posisi kredit debitur secara cepat.", "Identitas debitur, rekening, AO, status, saldo, dan detail kredit.", "User operasional"),
            ("mob", "MOB 6 Bulan", "Menganalisis kualitas pembayaran berdasarkan umur kredit enam bulan pertama.", "Cohort, MOB, bucket pembayaran, nominal, dan jumlah rekening.", "Manajemen risiko"),
            ("otp_baru", "Ontime Payment (OTP)", "Mengukur ketepatan pembayaran debitur terhadap tanggal jatuh tempo.", "Persentase OTP, nominal, NOA, periode, dan detail rekening.", "Collection / manajemen"),
            ("rekap_rr", "Repayment Rate (RR)", "Membandingkan posisi repayment closing bulan sebelumnya dengan actual harian berdasarkan flow migrasi.", "Nominal, NOA, persentase M-1, actual, migrasi, dan detail rekening terbatas.", "Manajemen / collection"),
            ("potensi_npl", "Potensi NPL", "Mengidentifikasi debitur yang berpotensi masuk NPL dan status penyelamatannya.", "Kandidat NPL, bucket, komitmen bayar, status penyelamatan, dan prioritas.", "Collection / remedial"),
            ("flow_par", "Flow PAR", "Memantau debitur yang berpindah ke kolektibilitas KL berdasarkan penyebab flow.", "Flow PAR, penyebab, nominal, NOA, kantor, AO, dan detail rekening.", "Manajemen risiko"),
            ("otp_bucket_fe", "OTP Bucket FE (31-90)", "Memantau ketepatan pembayaran untuk bucket yang menjadi fokus field collection.", "Bucket 31-90, nominal, NOA, pembayaran, dan prioritas follow-up.", "Field collection"),
        ],
    },
    {
        "name": "Collection dan Kualitas Aset",
        "audience": "Manajemen risiko, collection, dan remedial",
        "objective": "Menyajikan kualitas portofolio, migrasi kolektibilitas, bucket DPD, recovery, dan konsentrasi debitur bermasalah.",
        "reports": [
            ("npl", "Report NPL", "Menyajikan posisi NPL dan perbandingannya antara closing dengan actual.", "NPL bruto/neto, nominal, NOA, persentase, dan perubahan.", "Manajemen risiko"),
            ("migrasi_kolek", "Migrasi Kolek", "Mengikuti perpindahan kolektibilitas antarperiode.", "Matriks migrasi, nominal, NOA, arah pergerakan, dan detail.", "Manajemen risiko"),
            ("actual_kredit", "Bucket DPD & Kolek", "Membaca distribusi kualitas kredit menurut hari menunggak dan kolektibilitas.", "Bucket DPD, kolek L/DP/KL/D/M, saldo, NOA, dan proporsi.", "Collection / risiko"),
            ("migrasi_bucket", "Migrasi Bucket", "Menganalisis perpindahan debitur antar bucket pembayaran.", "Bucket asal, bucket tujuan, nominal, NOA, dan flow.", "Collection"),
            ("recovery_npl", "Recovery NPL", "Memantau recovery, posisi bersih, dan perubahan NPL dari closing ke actual.", "Recovery, saldo NPL, nominal, NOA, dan daftar rekening.", "Collection / remedial"),
            ("npl_25_besar", "25 NPL Besar", "Menyoroti 25 eksposur NPL terbesar untuk prioritas pengawasan.", "Debitur, saldo, kolek, AO, kantor, dan urutan eksposur.", "Pimpinan / remedial"),
            ("recovery_ph", "Hapus Buku", "Memantau rekening hapus buku dan perkembangan recovery pascahapus buku.", "Saldo hapus buku, recovery, status, dan detail rekening.", "Remedial"),
            ("maping_ao_remedial", "Mapping AO Remedial", "Memastikan rekening bermasalah memiliki penanggung jawab remedial yang jelas.", "Pemetaan debitur, AO remedial, kantor, bucket, dan status tindak lanjut.", "Remedial / pimpinan"),
        ],
    },
    {
        "name": "Laporan Keuangan dan Evaluasi Cabang",
        "audience": "Direksi, manajemen, dan pimpinan cabang",
        "objective": "Membaca posisi keuangan, pencapaian RBB, rasio, dan performa cabang dalam format ringkas maupun report evaluasi.",
        "reports": [
            ("lapkeu_kantor", "Laporan Keuangan", "Menyajikan laporan keuangan per kantor dengan tren dan komponen utama.", "Aset, kewajiban, pendapatan, beban, laba, tren, dan perbandingan kantor.", "Manajemen keuangan"),
            ("lap_neraca", "Lap Neraca", "Menampilkan posisi neraca aktual beserta struktur akun.", "Aset, liabilitas, ekuitas, kode akun, saldo, dan periode.", "Manajemen keuangan"),
            ("lap_laba_rugi", "Lap Laba Rugi", "Menampilkan pendapatan, beban, dan laba rugi aktual.", "Pendapatan, beban, laba rugi, saldo akun, dan periode.", "Manajemen keuangan"),
            ("rekap_lapkeu", "Rekap Lapkeu", "Meringkas laporan keuangan lintas kantor untuk perbandingan.", "Ringkasan aset, pendapatan, beban, laba, dan perubahan.", "Direksi / manajemen"),
            ("rbb_vs_realisasi", "RBB vs Realisasi", "Mengevaluasi pencapaian target RBB terhadap realisasi.", "Target, realisasi, selisih, capaian, dan sisa target.", "Manajemen bisnis"),
            ("ikhtisar", "Ikhtisar", "Menyajikan indikator utama bank dalam format ringkas dan seragam.", "KPMM/CAR, KAP, PPAP, NPL, kredit/aset produktif, ROA, NIM, BOPO, dan rasio lain.", "Direksi / manajemen"),
            ("realisasi_rbb_2026", "Realisasi RBB 2026", "Memantau realisasi RBB berjalan per wilayah dan kantor.", "RBB tahunan, target bulanan, realisasi, kekurangan/kelebihan, dan capaian.", "Manajemen wilayah"),
            ("paparan_rbb_realisasi", "Paparan RBB Direksi", "Menyediakan bahan paparan eksekutif untuk evaluasi produksi, keuangan, risiko, tren, dan SDM.", "Kinerja produksi, run off, ikhtisar, tren, RBB, SDM, dan ringkasan cabang.", "Direksi / manajemen"),
            ("raport_cabang", "Raport Cabang", "Mengevaluasi performa cabang dengan format report yang siap dibaca dan diekspor ke Excel.", "Keuangan, kredit, DPK, kualitas aset, RR, SDM, rasio, RBB, dan realisasi.", "Pimpinan cabang / atasan"),
            ("aging_kredit", "Rekap Aging Kredit", "Menganalisis usia kredit berdasarkan tanggal realisasi dan jatuh tempo.", "Umur kredit, tanggal realisasi, jatuh tempo, outstanding, dan kelompok aging.", "Manajemen kredit"),
        ],
    },
    {
        "name": "Layanan Digital",
        "audience": "Manajemen layanan dan pemilik produk digital",
        "objective": "Memantau pemanfaatan kanal digital dan aktivitas layanan yang mendukung transaksi serta akuisisi nasabah.",
        "reports": [
            ("layanan_digital", "Dashboard Layanan Digital", "Menyajikan ringkasan kinerja seluruh layanan digital.", "Tren transaksi, pengguna, volume, nilai transaksi, dan status layanan.", "Manajemen digital"),
            ("va", "Virtual Account (VA)", "Memantau penggunaan dan transaksi Virtual Account.", "Jumlah VA, transaksi, nominal, status, dan tren.", "Manajemen digital"),
            ("branchless", "Branchless", "Memantau aktivitas layanan branchless dan jangkauan layanan.", "Agen/outlet, aktivitas, transaksi, nominal, dan perkembangan.", "Manajemen digital"),
            ("qris_merchant", "QRIS Merchant", "Memantau pertumbuhan merchant dan transaksi QRIS.", "Merchant, transaksi, nominal, aktivitas, dan tren.", "Manajemen digital"),
        ],
    },
    {
        "name": "Dev Report dan Analisis Detail",
        "audience": "Divisi operasional, analis, dan pengelola laporan",
        "objective": "Menyediakan report operasional dengan filter, ringkasan, drill-down, dan detail rekening untuk analisis lanjutan.",
        "reports": [
            ("report_npl", "Report NPL", "Versi report terstruktur untuk eksplorasi NPL dan daftar pembentuknya.", "Filter kantor/AO, ringkasan NPL, rasio, dan detail rekening.", "Operasional / analis"),
            ("report_recovery_npl", "Report Recovery NPL", "Versi report terstruktur untuk mengevaluasi recovery NPL.", "Flow recovery, ringkasan, perbandingan, dan detail rekening.", "Operasional / remedial"),
            ("report_mutasi_kredit", "Report Mutasi Kredit", "Menyatukan realisasi, restrukturisasi, run off, dan perubahan portofolio.", "Mutasi nominal, arah perubahan, kantor, AO, dan detail rekening.", "Operasional / bisnis"),
            ("report_potensi_npl", "Report Potensi NPL", "Menyediakan daftar kandidat NPL yang dapat diprioritaskan.", "Ringkasan kandidat, status penyelamatan, komitmen bayar, dan detail.", "Operasional / collection"),
            ("report_flowpar", "Report Flow PAR", "Report analitik untuk menjelaskan pembentuk Flow PAR.", "Ringkasan penyebab, flow, nominal, NOA, dan detail rekening.", "Operasional / risiko"),
            ("rbb_produksi_kredit", "Report Produksi vs RBB", "Report produksi kredit yang memudahkan evaluasi capaian harian/bulanan.", "Kartu ringkasan, tabel bulanan, selisih, capaian, dan tren.", "Operasional / bisnis"),
            ("report_realisasi_ao", "Report Realisasi AO", "Membandingkan produksi, kelolaan, dan flow pada level AO.", "Kinerja AO, cohort kelolaan, flow, dan detail rekening.", "Operasional / pimpinan"),
            ("report_otp", "Report OTP", "Report ketepatan pembayaran dengan ringkasan dan detail pendukung.", "OTP, nominal, NOA, status pembayaran, dan daftar rekening.", "Operasional / collection"),
            ("pipelane_monitoring_kredit", "Monitoring Pipeline Kredit", "Memantau pipeline kredit lintas tahap sampai realisasi.", "Pipeline, tahapan, nominal, PIC/AO, aging proses, dan status.", "Operasional / bisnis"),
            ("prospek", "Pipeline Prospek", "Mencatat serta memonitor prospek yang berpotensi menjadi realisasi kredit.", "Prospek, sumber, PIC/AO, nilai, tahapan, dan tindak lanjut.", "Bisnis / AO"),
        ],
    },
    {
        "name": "Input dan Proyeksi RBB",
        "audience": "Pengelola RBB dan pemilik data unit kerja",
        "objective": "Menyiapkan target RBB bulanan sampai Desember sebagai sumber perhitungan pencapaian dan evaluasi.",
        "reports": [
            ("input_rbb", "Proyeksi RBB", "Merekap hasil input detail RBB per COA untuk Januari–Desember.", "COA, target bulanan, target tahunan, status draft/approval, dan rekap.", "Pengelola RBB"),
            ("input_rbb_aba", "Input RBB ABA", "Menginput penempatan ABA dan pendapatan bunga yang menjadi sumber proyeksi.", "Rekening ABA, nominal per bulan, pendapatan bunga, dan CKPN otomatis.", "Pengelola RBB"),
            ("input_rbb_detail?bagian=kredit", "Input RBB Kredit", "Menginput target RBB kredit per bulan sampai Desember.", "Target kredit bulanan, tahunan, sumber data, dan status input.", "Pengelola RBB kredit"),
            ("input_rbb_detail?bagian=damas", "Input RBB DAMAS", "Menginput target dana masyarakat/DAMAS per bulan sampai Desember.", "Target DAMAS bulanan, tahunan, dan rekap per komponen.", "Pengelola RBB dana"),
            ("input_rbb_detail?bagian=pendapatan", "Input RBB Pendapatan", "Menginput target pendapatan RBB per bulan sampai Desember.", "Target pendapatan, komponen COA, bulanan, tahunan, dan rekap.", "Pengelola RBB keuangan"),
            ("input_rbb_detail?bagian=beban", "Input RBB Beban", "Menginput target beban RBB per bulan sampai Desember.", "Target beban, komponen COA, bulanan, tahunan, dan rekap.", "Pengelola RBB keuangan"),
        ],
    },
    {
        "name": "KPI Bisnis",
        "audience": "Pimpinan, HR, dan pengelola kinerja",
        "objective": "Mengelola parameter, menghitung, menghasilkan, dan merekap penilaian KPI AO secara terukur.",
        "reports": [
            ("setting_kpi_jabatan", "Setting KPI Jabatan", "Mengatur indikator, bobot, arah penilaian, dan sumber data KPI.", "Parameter indikator, bobot, periode berlaku, dan status konfigurasi.", "Pengelola KPI"),
            ("hitung_kpi_ao", "Nilai KPI AO", "Menghitung nilai bulanan seorang AO sesuai jabatan dan indikator aktif.", "Nilai indikator, bobot, skor, hasil perhitungan, dan validasi.", "Pimpinan / HR"),
            ("generate_kpi_ao", "Generate KPI AO", "Memproses penilaian KPI AO untuk periode dan cabang yang dipilih.", "Status proses, hasil generate, jumlah AO, dan pesan validasi.", "Pengelola KPI"),
            ("rekap_kpi_ao", "Rekap KPI AO", "Merekap penilaian tahunan dari periode yang telah digenerate.", "Skor per AO, indikator, periode, peringkat, dan rekap tahunan.", "Pimpinan / HR"),
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:%s" % edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in kwargs[edge]:
                    element.set(qn("w:%s" % key), str(kwargs[edge][key]))


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:%s" % margin))
        if node is None:
            node = OxmlElement("w:%s" % margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_table(table, header_fill=BLUE, font_size=8.1, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r, row in enumerate(table.rows):
        if r == 0:
            repeat_header(row)
        for cell in row.cells:
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": BORDER}, bottom={"val": "single", "sz": 4, "color": BORDER}, left={"val": "single", "sz": 4, "color": BORDER}, right={"val": "single", "sz": 4, "color": BORDER})
            if r == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                        run.bold = True
                        run.font.size = Pt(font_size)
            elif r % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
    if widths:
        set_widths(table, widths)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_rule(paragraph, color=MID_BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    return p


def add_body(doc, text, *, bold_prefix=None, color=GRAY, size=9.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(BLUE)
        r1.font.name = "Arial"
        r1.font.size = Pt(size)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.color.rgb = RGBColor.from_string(color)
        r2.font.name = "Arial"
        r2.font.size = Pt(size)
    else:
        r = p.add_run(text)
        r.font.color.rgb = RGBColor.from_string(color)
        r.font.name = "Arial"
        r.font.size = Pt(size)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.04
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.2)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": 8, "color": accent}, bottom={"val": "single", "sz": 8, "color": accent}, left={"val": "single", "sz": 8, "color": accent}, right={"val": "single", "sz": 8, "color": accent})
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(GRAY)
    table.rows[0].cells[0].width = Inches(6.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_section_summary_table(doc, group):
    table = doc.add_table(rows=1, cols=5)
    headers = ["Menu", "Route", "Tujuan bisnis", "Output utama", "Pengguna"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color=WHITE, size=7.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for slug, name, purpose, output, audience in group["reports"]:
        row = table.add_row().cells
        set_cell_text(row[0], name, bold=True, color=NAVY, size=7.8)
        set_cell_text(row[1], f"pages/{slug}.php", color=GRAY, size=7.0)
        set_cell_text(row[2], purpose, color=GRAY, size=7.8)
        set_cell_text(row[3], output, color=GRAY, size=7.8)
        set_cell_text(row[4], audience, color=GRAY, size=7.2)
    style_table(table, font_size=7.8, widths=[1.18, 1.33, 1.66, 1.75, 0.78])
    return table


def add_catalog_table(doc):
    table = doc.add_table(rows=1, cols=5)
    headers = ["No.", "Area", "Menu report", "Route", "Peran utama"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color=WHITE, size=7.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    no = 1
    for group in REPORT_GROUPS:
        for slug, name, purpose, _output, audience in group["reports"]:
            row = table.add_row().cells
            set_cell_text(row[0], no, color=GRAY, size=7.6, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row[1], group["name"], color=BLUE, size=7.4)
            set_cell_text(row[2], name, bold=True, color=NAVY, size=7.7)
            set_cell_text(row[3], f"pages/{slug}.php", color=GRAY, size=6.8)
            set_cell_text(row[4], audience, color=GRAY, size=7.3)
            no += 1
    style_table(table, font_size=7.6, widths=[0.32, 1.28, 1.55, 1.55, 2.0])
    return table


def add_page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(GRAY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for level, size, color in [(1, 17, NAVY), (2, 13, BLUE), (3, 10.5, BLUE)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9 if level == 1 else 6)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.2)
        style.font.color.rgb = RGBColor.from_string(GRAY)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("MONBIS  |  DOKUMENTASI REPORT")
    hr.font.name = "Arial"
    hr.font.size = Pt(7.5)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor.from_string(BLUE)
    add_rule(hp, MID_BLUE)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Dokumentasi Report Monbis  •  Halaman ")
    fr.font.name = "Arial"
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = RGBColor.from_string(GRAY)
    add_field(fp, "PAGE")


def add_title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MONBIS")
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("DOKUMENTASI REPORT")
    r.font.name = "Arial"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Katalog menu dan panduan membaca laporan untuk manajemen")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_paragraph()
    add_callout(doc, "Tujuan dokumen", "Dokumen ini menjadi panduan ringkas bagi atasan untuk memahami laporan yang tersedia di Monbis, manfaatnya untuk evaluasi, serta output utama yang dapat digunakan sebagai bahan keputusan.", fill=LIGHT_BLUE, accent=BLUE)
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta_data = [("Produk", "Monitoring Bisnis (Monbis)"), ("Versi", "1.0 · Draft editable"), ("Tanggal", "30 September 2026"), ("Sasaran pembaca", "Direksi, manajemen, pimpinan cabang, dan pemilik proses")]
    for i, (left, right) in enumerate(meta_data):
        set_cell_text(meta.cell(i, 0), left, bold=True, color=BLUE, size=9.2)
        set_cell_text(meta.cell(i, 1), right, color=GRAY, size=9.2)
        set_cell_shading(meta.cell(i, 0), PALE_BLUE)
        for c in meta.rows[i].cells:
            set_cell_border(c, bottom={"val": "single", "sz": 4, "color": BORDER})
    set_widths(meta, [1.45, 5.0])
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dokumen ini dapat diedit kembali sesuai struktur organisasi, definisi indikator, dan kebijakan pelaporan yang berlaku.")
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)


def add_contents(doc):
    add_heading(doc, "Daftar Isi", 1)
    add_body(doc, "Gunakan judul bagian berikut sebagai peta baca. Struktur heading sudah disiapkan agar daftar isi Word dapat diperbarui kembali setelah dokumen diedit.", size=9.5)
    toc = doc.add_table(rows=1, cols=2)
    set_cell_text(toc.cell(0, 0), "Bagian", bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(toc.cell(0, 1), "Fokus", bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        ("1. Ringkasan untuk Atasan", "Cara menggunakan Monbis untuk evaluasi cepat dan keputusan."),
        ("2. Peta Report Monbis", "Daftar semua report aktif, route, dan peran utama."),
        ("3. Panduan per Area", "Tujuan dan output utama setiap kelompok menu."),
        ("4. Prioritas Baca Manajemen", "Urutan laporan untuk rapat harian, mingguan, dan bulanan."),
        ("5. Glosarium Indikator", "Istilah yang umum muncul di report Monbis."),
        ("6. Catatan Pengelolaan Dokumen", "Ruang edit, riwayat perubahan, dan aturan pemutakhiran."),
    ]
    for left, right in rows:
        r = toc.add_row().cells
        set_cell_text(r[0], left, bold=True, color=NAVY, size=8.7)
        set_cell_text(r[1], right, color=GRAY, size=8.7)
    style_table(toc, font_size=8.2, widths=[2.1, 4.35])


def add_exec_summary(doc):
    add_heading(doc, "1. Ringkasan untuk Atasan", 1)
    add_body(doc, "Monbis dirancang sebagai pusat monitoring bisnis yang menghubungkan kinerja produksi, pencapaian RBB, kualitas kredit, pembayaran, collection, keuangan, layanan digital, dan kinerja AO/cabang. Laporan dapat dibaca dari ringkasan eksekutif lalu dilanjutkan ke detail rekening atau unit kerja bila diperlukan.")
    add_callout(doc, "Cara membaca paling cepat", "Mulai dari Dashboard untuk melihat kondisi umum. Bila ada gap atau indikator merah, lanjutkan ke report area terkait: Produksi vs RBB untuk target, RR/OTP untuk pembayaran, NPL/Flow PAR untuk kualitas, dan Raport Cabang untuk evaluasi menyeluruh.", fill=GREEN, accent=TEAL)
    add_heading(doc, "Siklus evaluasi yang disarankan", 2)
    add_bullet(doc, "Harian: cek Dashboard, Produksi vs RBB, OTP/RR, Potensi NPL, dan alert yang muncul.")
    add_bullet(doc, "Mingguan: evaluasi produksi per AO/cabang, pipeline, jatuh tempo/refinancing, Flow PAR, migrasi bucket, dan recovery.")
    add_bullet(doc, "Bulanan: gunakan Raport Cabang, RBB vs Realisasi, Ikhtisar, laporan keuangan, NPL, dan KPI AO untuk rapat kinerja.")
    add_bullet(doc, "Tahunan: gunakan Realisasi RBB 2026, Paparan RBB Direksi, rekap KPI, dan tren keuangan sebagai bahan evaluasi target.")
    add_heading(doc, "Prinsip membaca angka", 2)
    add_bullet(doc, "Nominal menunjukkan skala dampak; persentase menunjukkan tingkat pencapaian atau perubahan; NOA menunjukkan jumlah rekening/debitur.")
    add_bullet(doc, "Perbandingan harus memperhatikan periode: bulan berjalan, bulan sebelumnya, akhir tahun sebelumnya, atau tahun berjalan.")
    add_bullet(doc, "Detail rekening dipakai untuk tindak lanjut, sedangkan ringkasan dipakai untuk keputusan dan prioritas.")


def add_catalog(doc):
    add_heading(doc, "2. Peta Report Monbis", 1)
    total = sum(len(g["reports"]) for g in REPORT_GROUPS)
    add_body(doc, f"Daftar berikut mencakup {total} menu report dan menu pendukung pelaporan yang tampil pada navigasi Monbis. Menu tertentu bersifat terbatas sesuai hak akses, khususnya Laporan, Dev Report, Input RBB, dan KPI Bisnis.")
    add_catalog_table(doc)
    add_body(doc, "Catatan: route dicantumkan agar dokumen mudah dipelihara oleh tim pengembang. Nama menu pada aplikasi dapat berubah mengikuti kebutuhan bisnis, namun tujuan dan output utamanya sebaiknya tetap diperbarui bersama.", size=8.2, color=GRAY, space_after=0)


def add_area_guides(doc):
    add_heading(doc, "3. Panduan per Area", 1)
    add_body(doc, "Bagian ini menjelaskan alasan keberadaan setiap kelompok laporan serta pertanyaan manajemen yang dapat dijawab oleh laporan tersebut.")
    for idx, group in enumerate(REPORT_GROUPS, start=1):
        add_heading(doc, f"3.{idx} {group['name']}", 2)
        add_body(doc, f"Fokus: {group['objective']}", bold_prefix="Fokus:")
        add_body(doc, f"Pengguna utama: {group['audience']}", bold_prefix="Pengguna utama:")
        add_section_summary_table(doc, group)
        doc.add_paragraph()
        if group["name"] == "Ringkasan Eksekutif":
            add_callout(doc, "Pertanyaan rapat", "Bagaimana posisi bisnis hari ini, indikator apa yang menyimpang, dan report mana yang harus dibuka untuk mengetahui penyebabnya?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Pemasaran":
            add_callout(doc, "Pertanyaan rapat", "Apakah produksi berjalan sesuai target, dari AO/cabang mana gap terbesar, dan berapa pipeline yang bisa dikonversi?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Monitoring Pembayaran":
            add_callout(doc, "Pertanyaan rapat", "Apakah penurunan RR/OTP berasal dari nominal besar, NOA besar, atau perpindahan flow tertentu yang perlu ditangani?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Collection dan Kualitas Aset":
            add_callout(doc, "Pertanyaan rapat", "Bucket atau kolektibilitas mana yang membentuk NPL, siapa penanggung jawabnya, dan berapa peluang recovery?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Laporan Keuangan dan Evaluasi Cabang":
            add_callout(doc, "Pertanyaan rapat", "Apakah pertumbuhan bisnis menghasilkan kinerja keuangan yang sehat dan bagaimana posisi cabang dibanding target serta periode pembanding?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Layanan Digital":
            add_callout(doc, "Pertanyaan rapat", "Kanal digital mana yang tumbuh, aktif digunakan, dan memberi kontribusi transaksi paling besar?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Dev Report dan Analisis Detail":
            add_callout(doc, "Pertanyaan rapat", "Detail rekening mana yang membentuk angka ringkasan dan tindakan apa yang perlu diteruskan ke unit kerja?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "Input dan Proyeksi RBB":
            add_callout(doc, "Pertanyaan rapat", "Apakah target sampai Desember sudah terisi lengkap, konsisten, dan siap menjadi dasar pengukuran pencapaian?", fill=LIGHT_BLUE, accent=BLUE)
        elif group["name"] == "KPI Bisnis":
            add_callout(doc, "Pertanyaan rapat", "Apakah indikator, bobot, dan hasil penilaian sudah sesuai kebijakan serta dapat dijelaskan kepada AO?", fill=LIGHT_BLUE, accent=BLUE)
        if idx != len(REPORT_GROUPS):
            add_page_break(doc)


def add_management_priority(doc):
    add_heading(doc, "4. Prioritas Baca Manajemen", 1)
    add_body(doc, "Urutan ini dapat digunakan sebagai pola rapat. Atasan tidak perlu membuka semua menu setiap saat; cukup mulai dari ringkasan lalu masuk ke area yang menunjukkan gap atau risiko.")
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Agenda", "Baca dulu", "Lanjutkan ke", "Output keputusan"]):
        set_cell_text(cell, text, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        ("Rapat harian", "Dashboard", "Produksi vs RBB; OTP/RR; Potensi NPL", "Prioritas tindak lanjut hari ini."),
        ("Rapat mingguan", "Realisasi Kredit / AO", "Pipeline; Jatuh Tempo; Flow PAR; Recovery", "Rencana aksi per AO/cabang."),
        ("Rapat bulanan", "Raport Cabang", "Ikhtisar; RBB vs Realisasi; Laporan Keuangan; NPL", "Evaluasi cabang dan koreksi target."),
        ("Rapat tahunan", "Realisasi RBB 2026", "Paparan RBB Direksi; KPI AO; tren keuangan", "Penetapan target dan arah bisnis."),
    ]
    for agenda, first, next_report, decision in rows:
        r = table.add_row().cells
        set_cell_text(r[0], agenda, bold=True, color=NAVY, size=8.4)
        set_cell_text(r[1], first, color=GRAY, size=8.4)
        set_cell_text(r[2], next_report, color=GRAY, size=8.4)
        set_cell_text(r[3], decision, color=GRAY, size=8.4)
    style_table(table, font_size=8.2, widths=[1.0, 1.45, 2.55, 1.45])
    add_heading(doc, "Paket laporan untuk bahan paparan", 2)
    add_bullet(doc, "Halaman pembuka: Dashboard + capaian utama RBB.")
    add_bullet(doc, "Halaman bisnis: Produksi vs RBB, Realisasi Kredit AO, Pipeline, dan Jatuh Tempo.")
    add_bullet(doc, "Halaman kualitas: RR, OTP, NPL, Flow PAR, Migrasi Kolek, dan Recovery.")
    add_bullet(doc, "Halaman cabang: Raport Cabang + Ikhtisar + laporan keuangan.")
    add_bullet(doc, "Halaman penggerak kinerja: KPI AO dan ringkasan SDM pada Paparan RBB Direksi.")


def add_glossary(doc):
    add_heading(doc, "5. Glosarium Indikator", 1)
    add_body(doc, "Definisi ringkas berikut membantu pembaca non-teknis memahami istilah yang umum muncul di laporan. Definisi detail/formula resmi dapat ditambahkan oleh pemilik proses sesuai kebijakan bank.")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Istilah", "Makna ringkas", "Dipakai untuk"]):
        set_cell_text(cell, text, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        ("RBB", "Rencana Bisnis Bank/target periode.", "Membandingkan target dan realisasi."),
        ("RR", "Repayment Rate, ukuran posisi pembayaran sesuai basis yang digunakan report.", "Mengevaluasi kemampuan bayar dan perubahan flow."),
        ("OTP", "Ontime Payment, ketepatan pembayaran terhadap jatuh tempo.", "Monitoring pembayaran harian/bulanan."),
        ("NPL", "Non-Performing Loan, kredit bermasalah sesuai definisi kualitas yang digunakan.", "Monitoring risiko kredit."),
        ("DPD", "Days Past Due/hari menunggak.", "Membentuk bucket keterlambatan."),
        ("Kolek", "Kolektibilitas kredit: L, DP, KL, D, dan M.", "Membaca kualitas rekening."),
        ("PAR", "Portfolio at Risk, portofolio dengan risiko keterlambatan sesuai threshold report.", "Flow dan monitoring kualitas."),
        ("NOA", "Number of Account, jumlah rekening/debitur.", "Melihat sebaran jumlah rekening selain nominal."),
        ("DPK/DAMAS", "Dana pihak ketiga/dana masyarakat.", "Monitoring tabungan dan deposito."),
        ("ROA/NIM/BOPO", "Rasio profitabilitas, margin, dan efisiensi operasional.", "Membaca kesehatan dan efisiensi keuangan."),
        ("Run Off", "Penurunan/outflow portofolio yang dibandingkan dengan realisasi baru.", "Menghitung kebutuhan realisasi dan pertumbuhan."),
        ("MOB", "Month on Book/umur kredit sejak realisasi.", "Membandingkan kualitas berdasarkan cohort umur."),
    ]
    for term, meaning, used in rows:
        r = table.add_row().cells
        set_cell_text(r[0], term, bold=True, color=NAVY, size=8.5)
        set_cell_text(r[1], meaning, color=GRAY, size=8.5)
        set_cell_text(r[2], used, color=GRAY, size=8.5)
    style_table(table, font_size=8.2, widths=[1.0, 3.25, 2.2])


def add_document_management(doc):
    add_heading(doc, "6. Catatan Pengelolaan Dokumen", 1)
    add_body(doc, "Bagian ini disiapkan agar dokumentasi dapat terus diperbarui saat ada penambahan menu, perubahan formula, perbaikan bug, atau perubahan hak akses.")
    add_callout(doc, "Aturan pemutakhiran", "Setiap perubahan yang berdampak pada angka, formula, filter, periode, tampilan, atau hak akses sebaiknya dicatat di riwayat berikut. Jika perubahan menambah menu, tambahkan juga satu baris pada Peta Report dan satu uraian pada Panduan per Area.", fill=ORANGE, accent="C55A11")
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ["Tanggal", "Menu/fitur", "Jenis perubahan", "Dampak bagi pengguna", "PIC / catatan"]):
        set_cell_text(cell, text, bold=True, color=WHITE, size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(8):
        r = table.add_row().cells
        for cell in r:
            set_cell_text(cell, "", color=GRAY, size=8.2)
            cell.height = Inches(0.26)
    style_table(table, font_size=8.0, widths=[0.78, 1.45, 1.45, 1.85, 1.02])
    add_heading(doc, "Template catatan perubahan", 2)
    add_bullet(doc, "Bug/perbaikan: tuliskan kondisi sebelum, kondisi sesudah, dan report yang terdampak.")
    add_bullet(doc, "Perubahan formula: tuliskan nama indikator, periode pembanding, sumber data, dan rumus ringkas.")
    add_bullet(doc, "Penambahan menu: tuliskan tujuan bisnis, pengguna, filter utama, output, dan route.")
    add_bullet(doc, "Perubahan hak akses: tuliskan role yang dapat melihat, mengubah, mengekspor, atau membuka detail.")
    add_heading(doc, "Catatan teknis untuk otomasi berikutnya", 2)
    add_body(doc, "Dokumen Word ini adalah baseline editable. Untuk tahap berikutnya, daftar report dapat dipindahkan ke katalog metadata terpusat sehingga perubahan menu dan deskripsi dapat digunakan sebagai sumber pembuatan Word/PDF secara otomatis pada setiap rilis.")


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.title = "Dokumentasi Report Monbis"
    core.subject = "Katalog laporan Monbis untuk manajemen"
    core.author = "Monbis"
    core.keywords = "Monbis, report, dokumentasi, manajemen, RBB, kredit"

    add_title_page(doc)
    add_page_break(doc)
    add_contents(doc)
    add_page_break(doc)
    add_exec_summary(doc)
    add_page_break(doc)
    add_catalog(doc)
    add_page_break(doc)
    add_area_guides(doc)
    add_page_break(doc)
    add_management_priority(doc)
    add_page_break(doc)
    add_glossary(doc)
    add_page_break(doc)
    add_document_management(doc)

    doc.save(OUTPUT)
    print(f"created {OUTPUT}")
    print(f"report_count {sum(len(g['reports']) for g in REPORT_GROUPS)}")


if __name__ == "__main__":
    build()
